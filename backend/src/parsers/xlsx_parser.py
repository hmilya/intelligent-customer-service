"""Excel .xlsx parser (each sheet as a labelled block)."""
from __future__ import annotations

from pathlib import Path

from .base import DocumentParseError, Parser


class XlsxParser(Parser):
    code = "xlsx"
    display_name = "Excel 表格"
    extensions = (".xlsx",)

    def parse(self, path: str | Path) -> str:
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as e:  # pragma: no cover
            raise DocumentParseError("openpyxl is not installed") from e

        p = Path(path)
        try:
            wb = load_workbook(filename=str(p), read_only=True, data_only=True)
        except Exception as e:
            raise DocumentParseError(f"Failed to open xlsx {p}: {e}") from e

        chunks: list[str] = []
        for ws in wb.worksheets:
            chunks.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [("" if v is None else str(v)).strip() for v in row]
                if any(cells):
                    chunks.append(" | ".join(cells))
            chunks.append("")
        wb.close()
        return "\n".join(chunks).strip()
