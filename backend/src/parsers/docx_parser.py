"""Word .docx parser (paragraphs + table cells)."""
from __future__ import annotations

from pathlib import Path

from .base import DocumentParseError, Parser


class DocxParser(Parser):
    code = "docx"
    display_name = "Word 文档"
    extensions = (".docx",)

    def parse(self, path: str | Path) -> str:
        try:
            from docx import Document  # type: ignore
        except ImportError as e:  # pragma: no cover
            raise DocumentParseError("python-docx is not installed") from e

        p = Path(path)
        try:
            doc = Document(str(p))
        except Exception as e:
            raise DocumentParseError(f"Failed to open docx {p}: {e}") from e

        chunks: list[str] = []
        for para in doc.paragraphs:
            text = (para.text or "").strip()
            if text:
                chunks.append(text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join((cell.text or "").strip() for cell in row.cells)
                if row_text.strip():
                    chunks.append(row_text)
        return "\n".join(chunks)
